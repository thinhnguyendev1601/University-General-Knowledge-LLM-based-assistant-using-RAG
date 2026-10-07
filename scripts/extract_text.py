#!/usr/bin/env python3
"""Turn raw documents into CANONICAL TEXT (the reference every offset points into).

    data/raw/<doc_id>.pdf|.docx   ->   data/processed/<doc_id>.txt
                                       data/processed/<doc_id>.pages.json

Usage:
    python scripts/extract_text.py                 # all files in data/raw
    python scripts/extract_text.py moet-tt08-2021  # one doc_id

This is a deliberately plain v0 extractor so annotation can start today. The ingestion
owner may later build a better parser (tables, layout) — but the resulting canonical text
of a document version is FROZEN once test questions cite it (evidence quotes are resolved
against it). If you must re-extract, bump doc_version and re-run `eval.testset build`.

Requires: pip install pymupdf python-docx
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from core.textnorm import normalize_text, sha256_text, text_quality  # noqa: E402

RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "processed"
PAGE_SEP = "\n\n"


def extract_pdf(path: Path) -> list[str]:
    import pymupdf  # PyMuPDF

    pages = []
    with pymupdf.open(path) as doc:
        for page in doc:
            pages.append(page.get_text("text", sort=True))
    return pages


def extract_docx(path: Path) -> list[str]:
    import docx  # python-docx

    d = docx.Document(str(path))
    parts: list[str] = []
    for p in d.paragraphs:
        parts.append(p.text)
    for t in d.tables:  # tables appended after body in v0 (order not preserved)
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return ["\n".join(parts)]  # DOCX has no stable pages -> one "page"


def process(path: Path) -> dict:
    doc_id = path.stem
    ext = path.suffix.lower()
    if ext == ".pdf":
        raw_pages, parser = extract_pdf(path), "pymupdf"
    elif ext == ".docx":
        raw_pages, parser = extract_docx(path), "python-docx"
    else:
        raise ValueError(f"unsupported format: {path.name}")

    pages = [normalize_text(p) for p in raw_pages]
    text_parts, page_starts, pos = [], [], 0
    for i, p in enumerate(pages):
        page_starts.append(pos)
        text_parts.append(p)
        pos += len(p) + (len(PAGE_SEP) if i < len(pages) - 1 else 0)
    text = PAGE_SEP.join(text_parts)

    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / f"{doc_id}.txt").write_text(text, encoding="utf-8")
    meta = {
        "doc_id": doc_id,
        "parser": parser,
        "n_pages": len(pages),
        "page_starts": page_starts,  # char offset where each page begins (0-based list, page = idx+1)
        "page_separator": PAGE_SEP,
        "text_sha256": sha256_text(text),
        "quality": text_quality(text, len(pages)),
    }
    (OUT / f"{doc_id}.pages.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return meta


def warn(meta: dict) -> list[str]:
    q, w = meta["quality"], []
    if q["chars_per_page"] < 300:
        w.append("very little text per page -> scanned/image PDF? (needs OCR or exclude from v0)")
    if q["replacement_chars"] > 0:
        w.append(f"{q['replacement_chars']} replacement chars (U+FFFD) -> encoding problem")
    if q["non_letter_ratio"] > 0.35:
        w.append("high non-letter ratio -> possible garbled text (legacy VN fonts?)")
    return w


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("doc_ids", nargs="*")
    args = ap.parse_args()

    files = sorted(p for p in RAW.glob("*") if p.suffix.lower() in {".pdf", ".docx"})
    if args.doc_ids:
        files = [p for p in files if p.stem in set(args.doc_ids)]
    if not files:
        print(f"no input files found in {RAW}")
        return
    print(f"{'doc_id':32} {'pages':>5} {'chars':>8} {'ch/pg':>7} {'vi-diac':>8}  warnings")
    for p in files:
        try:
            m = process(p)
        except Exception as e:  # keep going, report at the end
            print(f"{p.stem:32} FAILED: {e}")
            continue
        q = m["quality"]
        w = warn(m)
        print(f"{p.stem:32} {m['n_pages']:>5} {q['n_chars']:>8} {q['chars_per_page']:>7} {q['vi_diacritic_ratio']:>8}  {'; '.join(w) or 'ok'}")


if __name__ == "__main__":
    main()
